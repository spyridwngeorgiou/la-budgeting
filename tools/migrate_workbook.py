"""
migrate_workbook.py -- Επιχειρησιακό_Αρχείο_107.xlsx -> seed SQL for backoffice/

Reads the hand-kept source-of-truth workbook and writes a REVIEWABLE .sql
file (never live inserts) using ON CONFLICT (natural key) DO UPDATE, so the
script is safely rerunnable. Fails loudly -- raises UnmappedLiteralError --
on any Greek literal it doesn't recognise, rather than silently defaulting.

Usage:
    py tools/migrate_workbook.py --dry-run
    py tools/migrate_workbook.py --out seed/0100_migrated_data.sql

Scope of this pass: orgs/settings, projects, contacts, accounts, categories,
and all 102 one-off Κινήσεις rows from the hand-kept workbook.

Deliberately NOT yet handled here (see tools/MIGRATION_NOTES.md):
  - The 11 recurring/installment plans, which live only in the *generated*
    back_office_operations.xlsx (Table_Kin columns S/U and the per-row
    installment fields), not in the hand-kept source of truth.
  - Recovering true invoice numbers by joining aade_exports/AADE_Master.xlsx
    on ΜΑΡΚ (today's counterparty_afm/invoice_number come from the workbook's
    own VLOOKUP, which is blank for contacts missing from Επαφές).
  - project_budgets / project_model_inputs from sheet "Προϋπολογισμοί".
  - assets / liabilities / expected_income from the Λογαριασμοί wealth blocks.
Each is a natural next increment on top of the masters + ledger this script
lands, using the same fail-loud mapping approach.
"""

from __future__ import annotations

import argparse
import re
import sys
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from mappings import (  # noqa: E402
    ACCOUNT_KIND,
    DIRECTION,
    OWNER_SCOPE,
    ORIGIN,
    SCOPE,
    STATUS,
    UnmappedLiteralError,
    map_literal,
)

ROOT = Path(__file__).parent.parent
SOURCE_WORKBOOK = ROOT / "Επιχειρησιακό_Αρχείο_107.xlsx"

# Deterministic UUIDs (uuid5 over a fixed namespace + natural key) so the
# generated SQL is byte-stable across reruns and diffable in code review.
NAMESPACE = uuid.UUID("6f6e0a4e-2f61-4e1a-9c1a-9f6b6f9a5b10")

# Strict project-code shape (Q000_GENERAL, Q001_ILIOUPOLI_P15_RESIDENCES, ...).
# A loose startswith("Q0") check false-positives on note text like
# "Q003 ΛΑΖΑΡΑΚΗ 32" that appears elsewhere in the same sheet's pick-lists.
PROJECT_CODE_RE = re.compile(r"^Q\d{3}_[A-Z0-9_]+$")


def stable_id(*parts: str) -> str:
    return str(uuid.uuid5(NAMESPACE, "|".join(parts)))


def sql_str(value) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def sql_num(value) -> str:
    if value is None:
        return "null"
    return str(value)


def sql_date(value) -> str:
    if value is None:
        return "null"
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return f"'{value.isoformat()}'"
    return sql_str(value)


def sql_bool(value: bool | None) -> str:
    if value is None:
        return "null"
    return "true" if value else "false"


def normalize_afm(value) -> str | None:
    if value is None:
        return None
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    if not digits:
        return None
    return digits.zfill(9)


def normalize_name(value) -> str:
    return " ".join(str(value or "").split()).upper()


@dataclass
class MigrationState:
    org_id: str
    project_ids: dict[str, str] = field(default_factory=dict)   # code -> id
    contact_ids: dict[str, str] = field(default_factory=dict)   # afm or name -> id
    contact_afms: dict[str, str] = field(default_factory=dict)  # contact id -> afm
    account_ids: dict[str, str] = field(default_factory=dict)   # name -> id
    category_ids: dict[str, str] = field(default_factory=dict)  # name -> id
    unmapped: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    row_counts: dict[str, int] = field(default_factory=dict)


def load_workbook():
    if not SOURCE_WORKBOOK.exists():
        raise FileNotFoundError(f"Source workbook not found: {SOURCE_WORKBOOK}")
    return openpyxl.load_workbook(SOURCE_WORKBOOK, data_only=True)


def emit_org(state: MigrationState, org_afm: str, org_name: str) -> str:
    return f"""
insert into orgs (id, name, own_afm)
values ({sql_str(state.org_id)}, {sql_str(org_name)}, {sql_str(org_afm)})
on conflict (id) do update set name = excluded.name, own_afm = excluded.own_afm;
"""


def emit_projects(state: MigrationState, wb) -> list[str]:
    ws = wb["Ρυθμίσεις"]
    codes: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        for cell in row:
            if isinstance(cell, str) and PROJECT_CODE_RE.match(cell.strip()):
                codes.append(cell.strip())
    codes = sorted(set(codes))

    display_names = {
        "Q000_GENERAL": "Γενικά εταιρείας",
        "Q001_ILIOUPOLI_P15_RESIDENCES": "Ηλιούπολη",
        "Q002_VOULIAGMENI_309": "Βουλιαγμένη",
        "Q003_LAZARAKI32_GLYFADA": "Λαζαράκη 32, Γλυφάδα",
        "Q004_AGIOU_KWNSTANTINOU20_GLYFADA": "Αγ. Κωνσταντίνου 20, Γλυφάδα",
        "Q005_LEGRENA": "Λεγρενά",
        "Q006_KAVOURI_AKTIS2": "Καβούρι, Ακτής 2",
        "Q007_KOLONAKI_KARNEADOU37": "Κολωνάκι, Καρνεάδου 37",
    }

    stmts = []
    for i, code in enumerate(codes):
        pid = stable_id("project", code)
        state.project_ids[code] = pid
        name = display_names.get(code, code)
        stmts.append(f"""
insert into projects (id, org_id, code, display_name, sort_order)
values ({sql_str(pid)}, {sql_str(state.org_id)}, {sql_str(code)}, {sql_str(name)}, {i})
on conflict (org_id, code) do update set display_name = excluded.display_name;
""")
    state.row_counts["projects"] = len(codes)
    return stmts


def emit_contacts(state: MigrationState, wb) -> list[str]:
    ws = wb["Επαφές"]
    stmts = []
    n = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        name = row[0]
        if not name or not str(name).strip() or str(name).strip() == "ΣΥΝΟΛΟ":
            continue
        afm = normalize_afm(row[1]) if len(row) > 1 else None
        key = afm or normalize_name(name)
        cid = stable_id("contact", key)
        state.contact_ids[key] = cid
        if afm:
            state.contact_afms[cid] = afm
        # Index by name too, so ledger rows can resolve contacts that only
        # give a name (many Κινήσεις rows have no ΑΦΜ column filled in).
        state.contact_ids[normalize_name(name)] = cid

        stmts.append(f"""
insert into contacts (id, org_id, name, afm)
values ({sql_str(cid)}, {sql_str(state.org_id)}, {sql_str(str(name).strip())}, {sql_str(afm)})
on conflict (org_id, afm) where afm is not null and afm <> '' do update
  set name = excluded.name;
""")
        n += 1
    state.row_counts["contacts"] = n
    return stmts


def emit_accounts(state: MigrationState, wb) -> list[str]:
    ws = wb["Λογαριασμοί"]
    stmts = []
    n = 0
    balances_date = date(2026, 8, 21)  # from Ρυθμίσεις "Ημερομηνία Υπολοίπων Έναρξης"

    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(
        (i for i, r in enumerate(rows) if r and str(r[0] or "").strip() == "Λογαριασμός"),
        None,
    )
    if header_idx is None:
        raise RuntimeError("Could not find 'Λογαριασμός' header in Λογαριασμοί sheet")

    # Only the true liquid cash/bank rows live under this header (columns:
    # name, opening_balance, movement, current_balance, owner_scope, notes).
    # Χρυσός/Κρυπτονομίσματα appear further down under a DIFFERENT header
    # ("Στοιχείο / Εκτιμώμενη Αξία / Ποσοστό") that is really the shape of
    # the `assets` table (estimated value + ownership %), not an account with
    # an opening balance -- they belong in tools/migrate_workbook.py's
    # not-yet-built assets step, alongside the real estate and watches they
    # sit next to in the sheet. Stop at the first "ΣΥΝΟΛΟ" or blank row.
    for row in rows[header_idx + 1:]:
        name = row[0]
        if name is None or not str(name).strip() or str(name).strip().startswith("ΣΥΝΟΛΟ"):
            break
        opening = row[1] if row[1] is not None else 0
        owner_literal = row[4] if len(row) > 4 else None
        try:
            owner_scope = map_literal(OWNER_SCOPE, owner_literal, column="Κάτοχος")
        except UnmappedLiteralError as e:
            state.unmapped.append(str(e))
            continue

        kind = "cash" if "Μετρητά" in str(name) else "bank"

        aid = stable_id("account", str(name).strip())
        state.account_ids[str(name).strip()] = aid
        stmts.append(f"""
insert into accounts (id, org_id, name, kind, owner_scope, is_liquid, opening_balance, opening_balance_date)
values ({sql_str(aid)}, {sql_str(state.org_id)}, {sql_str(str(name).strip())}, {sql_str(kind)},
        {sql_str(owner_scope)}, true, {sql_num(opening)}, {sql_date(balances_date)})
on conflict (org_id, name) do update
  set opening_balance = excluded.opening_balance, opening_balance_date = excluded.opening_balance_date;
""")
        n += 1
    state.row_counts["accounts"] = n
    return stmts


def emit_categories(state: MigrationState, wb) -> list[str]:
    ws = wb["Ρυθμίσεις"]
    header = None
    col_idx = None
    for row in ws.iter_rows(values_only=True):
        if row and "Κατηγορία" in row:
            header = row
            col_idx = row.index("Κατηγορία")
            break
    if header is None:
        raise RuntimeError("Could not find 'Κατηγορία' column in Ρυθμίσεις sheet")

    stmts = []
    n = 0
    seen = set()
    rows = list(ws.iter_rows(values_only=True))
    start = rows.index(header) + 1
    for row in rows[start:]:
        val = row[col_idx] if col_idx < len(row) else None
        # The pick-list is a contiguous run right below the header; the same
        # column is reused further down for an unrelated assumptions block
        # (rates, dates, business-plan figures). Stop at the first blank
        # cell rather than skipping blanks and scanning past the list --
        # that scan is what previously picked up stray strings like 'Ναι'
        # and '2026-11' from the reused block.
        if val is None or (isinstance(val, str) and not val.strip()):
            break
        if not isinstance(val, str):
            continue
        name = val.strip()
        if name in seen:
            continue
        seen.add(name)
        code = name.upper().replace(" ", "_").replace("&", "AND")
        cat_id = stable_id("category", code)
        state.category_ids[name] = cat_id
        is_financing = name == "Δάνειο"
        stmts.append(f"""
insert into categories (id, org_id, code, name, is_financing)
values ({sql_str(cat_id)}, {sql_str(state.org_id)}, {sql_str(code)}, {sql_str(name)}, {sql_bool(is_financing)})
on conflict (org_id, code) do update set name = excluded.name;
""")
        n += 1
    state.row_counts["categories"] = n
    return stmts


def resolve_contact(state: MigrationState, name, afm) -> str | None:
    norm_afm = normalize_afm(afm)
    if norm_afm and norm_afm in state.contact_ids:
        return state.contact_ids[norm_afm]
    if name:
        norm_name = normalize_name(name)
        if norm_name in state.contact_ids:
            return state.contact_ids[norm_name]
    return None


def emit_transactions(state: MigrationState, wb) -> list[str]:
    ws = wb["Κινήσεις"]
    stmts = []
    n = 0
    for row_no, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        if row[0] in (None, ""):
            continue

        (tx_date, paid_date, due_date, contact_name, afm, project_code,
         category_name, description, tx_type_lit, net, vat, withholding,
         gross, account_name, status_lit, origin_lit) = row[:16]

        try:
            direction = map_literal(DIRECTION, tx_type_lit, column="Τύπος", row=row_no)
            status = map_literal(STATUS, status_lit, column="Κατάσταση", row=row_no)
            origin = map_literal(ORIGIN, origin_lit, column="Προέλευση", row=row_no) or "manual"
        except UnmappedLiteralError as e:
            state.unmapped.append(str(e))
            continue

        # origin='aade' carries a real-world promise: the DB's fingerprint
        # uniqueness guard (tx_mark_uq / tx_fingerprint_uq, scoped to
        # origin='aade') exists specifically to stop re-imports of the same
        # myDATA invoice from double-inserting. This migration hasn't yet
        # joined AADE_Master.xlsx to recover real ΜΑΡΚ values (see
        # MIGRATION_NOTES.md #2), so a workbook row merely LABELLED "AADE"
        # has no mark to disambiguate it from a same-day, same-amount,
        # same-contact row that is a genuinely different invoice (this
        # happens here: a credit note and its same-day re-invoice). Downgrade
        # to 'manual' until the mark backfill lands and can re-classify these
        # with real provenance.
        if origin == "aade":
            origin = "manual"

        project_id = state.project_ids.get(str(project_code).strip()) if project_code else None
        if project_code and project_id is None:
            state.unmapped.append(
                f"Unknown project code {project_code!r} at Κινήσεις row {row_no}"
            )
            continue

        contact_id = resolve_contact(state, contact_name, afm)
        account_id = state.account_ids.get(str(account_name).strip()) if account_name else None
        category_id = state.category_ids.get(str(category_name).strip()) if category_name else None

        net = net if net is not None else 0
        vat = vat if vat is not None else 0
        withholding = withholding if withholding is not None else 0
        gross_computed = round(float(net) + float(vat) - float(withholding), 2)

        # A handful of AADE rows are credit notes (Πιστωτικό/Αντιλογισμός)
        # carrying negative net/VAT -- e.g. row 43, SPECTER DESIGN GROUP,
        # -1275/-306. The schema enforces "amounts always positive, direction
        # carries the sign" (matching the workbook's own QC rule), so a
        # credit note received against an expense is migrated as the
        # opposite direction with absolute values -- economically it *is*
        # money coming back, same as income. Flagged for human review since
        # this reclassification is a judgment call, not a pure transcription.
        if gross_computed < 0:
            direction = "income" if direction == "expense" else "expense"
            net, vat, withholding = -net, -vat, -withholding
            gross_computed = -gross_computed
            state.notes.append(
                f"row {row_no} had a negative amount ({contact_name!r}, "
                f"{description!r}) -- migrated as '{direction}' with absolute "
                f"values. Review before trusting."
            )

        # The workbook's own ΑΦΜ column on Κινήσεις is a VLOOKUP into Επαφές
        # and is blank on many rows (a known live dedup hole, per the plan).
        # Backfill from the resolved contact's own ΑΦΜ when the row itself
        # didn't carry one, so counterparty_afm is populated whenever
        # possible rather than only when the workbook happened to show it.
        counterparty_afm = normalize_afm(afm) or (
            state.contact_afms.get(contact_id) if contact_id else None
        )
        has_invoice = bool(vat) or bool(withholding)

        tx_id = stable_id("tx", "kiniseis-row", str(row_no))
        # Keyed on the workbook's own row number, not on content: two
        # genuinely distinct transactions can share date+afm+amount (blank
        # AFM plus a round number recurs often here), which a content key
        # would collide on and silently drop. Row position is what actually
        # needs to be idempotent across reruns of the same workbook version.
        # This is deliberately separate from `fingerprint` (recomputed in the
        # DB from date+afm+gross+mark), which serves AADE re-import dedup,
        # not migration idempotency.
        legacy_key = f"kiniseis-row-{row_no}"

        stmts.append(f"""
insert into transactions (
  id, org_id, tx_date, paid_on, due_date, contact_id, counterparty_afm, counterparty_name,
  project_id, category_id, account_id, direction, scope, status, origin,
  gross_amount, net_amount, vat_amount, withholding_amount, has_invoice,
  description, legacy_excel_id
) values (
  {sql_str(tx_id)}, {sql_str(state.org_id)}, {sql_date(tx_date)}, {sql_date(paid_date)}, {sql_date(due_date)},
  {sql_str(contact_id)}, {sql_str(counterparty_afm)}, {sql_str(contact_name)},
  {sql_str(project_id)}, {sql_str(category_id)}, {sql_str(account_id)}, {sql_str(direction)},
  'business', {sql_str(status)}, {sql_str(origin)},
  {sql_num(gross_computed)}, {sql_num(net)}, {sql_num(vat)}, {sql_num(withholding)}, {sql_bool(has_invoice)},
  {sql_str(description)}, {sql_str(legacy_key)}
)
on conflict (org_id, legacy_excel_id) where legacy_excel_id is not null do update set
  gross_amount = excluded.gross_amount, net_amount = excluded.net_amount,
  vat_amount = excluded.vat_amount, withholding_amount = excluded.withholding_amount;
""")
        n += 1
    state.row_counts["transactions"] = n
    return stmts


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--org-afm", default="000000000",
                         help="Your company's ΑΦΜ -- required to determine AADE import direction")
    args = parser.parse_args()

    wb = load_workbook()
    state = MigrationState(org_id=stable_id("org", "default"))

    all_stmts = []
    all_stmts.append(emit_org(state, args.org_afm, "Η επιχείρησή μου"))
    all_stmts += emit_projects(state, wb)
    all_stmts += emit_contacts(state, wb)
    all_stmts += emit_accounts(state, wb)
    all_stmts += emit_categories(state, wb)
    all_stmts += emit_transactions(state, wb)

    print("Row counts:")
    for entity, count in state.row_counts.items():
        print(f"  {entity:14s} {count}")

    if state.unmapped:
        print(f"\n{len(state.unmapped)} unmapped literal(s) -- these rows were SKIPPED:")
        for msg in state.unmapped:
            print(f"  - {msg}")

    if state.notes:
        print(f"\n{len(state.notes)} note(s) -- not errors, but please verify:")
        for msg in state.notes:
            print(f"  - {msg}")

    if args.dry_run:
        if state.unmapped:
            sys.exit(1)
        print("\nDry run OK, no unmapped literals.")
        return

    if state.unmapped:
        print("\nRefusing to write seed file while unmapped literals remain.")
        sys.exit(1)

    out_path = args.out or (ROOT / "backoffice" / "seed" / "0100_migrated_data.sql")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as f:
        f.write("-- Generated by tools/migrate_workbook.py -- review before applying.\n")
        f.write("begin;\n")
        f.writelines(all_stmts)
        f.write("commit;\n")
    print(f"\nWrote {out_path}")


if __name__ == "__main__":
    main()
